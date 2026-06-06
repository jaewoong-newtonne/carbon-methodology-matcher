"""Build the GS methodology URL index from the eligibility xlsx.

Source: the eligible-CDM-GS-methodologies xlsx (see --xlsx)
Target: `carbonmm/configs/gs_methodology_index.yaml`

The xlsx is maintained by Gold Standard and lists the eligible methodologies
(GS-native + CDM-approved-for-GS). Here we need only the 'Gold Standard meths'
sheet — CDM methodologies are covered by the published CDM methodology booklet.

This script:
  1. Opens the xlsx via openpyxl
  2. Iterates the 'Gold Standard meths' sheet, extracts hyperlinked cells
  3. Parses each hyperlink URL into (code_prefix, slug)
  4. Writes a YAML index suitable for scrape_gs.py --resolve

Rerun whenever the xlsx is updated (Gold Standard publishes new versions
periodically — current is V2.11 dated 2025-04-16).

Usage:
    python -m carbonmm.ingest.build_gs_index
    python -m carbonmm.ingest.build_gs_index --xlsx path/to/file.xlsx
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

logger = logging.getLogger(__name__)

DEFAULT_XLSX = Path(__file__).resolve().parents[3] / "data" / "reference" / "427_V2.11_List-of-eligible-CDM-GS-methodologies.xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "configs" / "gs_methodology_index.yaml"
TARGET_SHEET = "Gold Standard meths"

GS_URL_PATTERN = re.compile(r"https://globalgoals\.goldstandard\.org/([^/]+)/?")
CODE_PREFIX_PATTERN = re.compile(r"^(\d+(?:[-.]\d+)*)-")


def parse_slug(url: str) -> tuple[str | None, str | None]:
    """Extract (slug, code_prefix) from a globalgoals.goldstandard.org URL."""
    m = GS_URL_PATTERN.match(url)
    if not m:
        return None, None
    slug = m.group(1)
    cm = CODE_PREFIX_PATTERN.match(slug)
    code = cm.group(1).replace(".", "-") if cm else None
    return slug, code


def build_index(xlsx_path: Path, sheet_name: str = TARGET_SHEET) -> list[dict]:
    """Walk the GS sheet and collect entries with hyperlinks."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not in {wb.sheetnames}")
    ws = wb[sheet_name]

    entries = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink and cell.hyperlink.target and cell.hyperlink.target.startswith("http"):
                slug, code = parse_slug(cell.hyperlink.target)
                entries.append({
                    "row": cell.row,
                    "code_prefix": code,
                    "slug": slug,
                    "name": str(cell.value or "").strip(),
                    "url": cell.hyperlink.target,
                })
    return entries


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="Source xlsx (default: %(default)s)")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Target YAML (default: %(default)s)")
    parser.add_argument("--sheet", default=TARGET_SHEET)
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(name)-25s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
    )

    if not args.xlsx.exists():
        logger.error("xlsx not found: %s", args.xlsx)
        return 1

    entries = build_index(args.xlsx, args.sheet)
    logger.info("Extracted %d entries from %s", len(entries), args.sheet)

    import yaml
    args.output.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": "1.0",
        "source": str(args.xlsx.relative_to(Path(__file__).resolve().parents[3])),
        "sheet": args.sheet,
        "extracted_at": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "count": len(entries),
        "note": (
            "GS-native methodologies (registry=GS) eligible for Gold Standard "
            "projects. CDM methodologies eligible-for-GS are covered separately "
            "by the published CDM methodology booklet — no web scraping needed "
            "for CDM. Regenerate this YAML via build_gs_index.py when the source "
            "xlsx is updated."
        ),
        "entries": entries,
    }
    with open(args.output, "w") as f:
        yaml.safe_dump(payload, f, sort_keys=False, allow_unicode=True, default_flow_style=False)
    logger.info("Wrote → %s", args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
